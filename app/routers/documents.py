from fastapi import (
    APIRouter,
    UploadFile,
    File,
    Form,
    HTTPException
)

from pydantic import BaseModel
from app.database import db

from uuid import uuid4
from pathlib import Path

import os
import csv
import openpyxl


router = APIRouter(
    prefix="/documents",
    tags=["Documents"]
)


# =====================================================
# STORAGE
# =====================================================

# Project root:
# backend/
# ├── app/
# ├── storage/
# └── venv/

BASE_DIR = Path(
    __file__
).resolve().parent.parent

UPLOAD_DIR = BASE_DIR / "storage"

UPLOAD_DIR.mkdir(
    parents=True,
    exist_ok=True
)

print("DOCUMENT STORAGE:", UPLOAD_DIR)


# =====================================================
# UPLOAD DOCUMENT
# =====================================================

@router.post("/upload")
async def upload_document(
    workspace_id: str = Form(...),
    file: UploadFile = File(...)
):

    if not workspace_id:
        raise HTTPException(
            status_code=400,
            detail="workspace_id is required"
        )

    if not file.filename:
        raise HTTPException(
            status_code=400,
            detail="File is required"
        )

    # Check workspace
    workspace = db.workspaces.find_one(
        {"id": workspace_id}
    )

    if workspace is None:
        raise HTTPException(
            status_code=404,
            detail="Workspace not found"
        )

    # Generate document ID
    document_id = str(uuid4())

    # Safe filename
    filename = Path(
        file.filename
    ).name

    # Actual physical file path
    physical_path = (
        UPLOAD_DIR /
        f"{document_id}_{filename}"
    )

    # Read uploaded file
    contents = await file.read()

    if not contents:
        raise HTTPException(
            status_code=400,
            detail="Uploaded file is empty"
        )

    # Save physical file
    try:

        with open(
            physical_path,
            "wb"
        ) as f:

            f.write(contents)

    except Exception as e:

        raise HTTPException(
            status_code=500,
            detail=f"Could not save file: {str(e)}"
        )

    # Verify file actually exists
    if not physical_path.is_file():

        raise HTTPException(
            status_code=500,
            detail="File was not saved correctly"
        )

    # Save absolute path in MongoDB
    document_data = {

        "id": document_id,

        "workspace_id": workspace_id,

        "filename": filename,

        "file_type": file.content_type,

        "file_path": str(
            physical_path
        ),

        "status": "uploaded"

    }

    db.documents.insert_one(
        document_data
    )

    print("================================")
    print("DOCUMENT UPLOADED")
    print("ID:", document_id)
    print("Filename:", filename)
    print("Path:", physical_path)
    print("Exists:", physical_path.is_file())
    print("================================")

    return {

        "message":
            "Document uploaded successfully",

        "document": {

            "id": document_id,

            "workspace_id":
                workspace_id,

            "filename":
                filename,

            "file_type":
                file.content_type,

            "file_path":
                str(physical_path),

            "status":
                "uploaded"

        }

    }


# =====================================================
# GET DOCUMENTS BY WORKSPACE
# =====================================================

@router.get("/workspace/{workspace_id}")
def get_documents(
    workspace_id: str
):

    documents = list(
        db.documents.find(
            {
                "workspace_id":
                    workspace_id
            },
            {
                "_id": 0
            }
        )
    )

    return {

        "workspace_id":
            workspace_id,

        "count":
            len(documents),

        "documents":
            documents

    }


# =====================================================
# GET SINGLE DOCUMENT
# =====================================================

@router.get("/document/{document_id}")
def get_document(
    document_id: str
):

    document = db.documents.find_one(
        {
            "id": document_id
        },
        {
            "_id": 0
        }
    )

    if document is None:

        raise HTTPException(
            status_code=404,
            detail="Document not found"
        )

    return {
        "document": document
    }


# =====================================================
# PREVIEW CSV / XLSX
# =====================================================

@router.get("/preview/{document_id}")
def preview_document(
    document_id: str
):

    document = db.documents.find_one(
        {
            "id": document_id
        },
        {
            "_id": 0
        }
    )

    if document is None:

        raise HTTPException(
            status_code=404,
            detail="Document not found"
        )

    file_path = document.get(
        "file_path"
    )

    if not file_path:

        raise HTTPException(
            status_code=404,
            detail="File path not found"
        )

    file_path = Path(
        file_path
    )

    if not file_path.is_file():

        raise HTTPException(
            status_code=404,
            detail=f"File not found: {file_path}"
        )

    filename = document[
        "filename"
    ].lower()


    # CSV
    if filename.endswith(".csv"):

        rows = []

        with open(
            file_path,
            "r",
            encoding="utf-8-sig",
            newline=""
        ) as csv_file:

            reader = csv.reader(
                csv_file
            )

            for row in reader:
                rows.append(row)

        return {

            "document_id":
                document_id,

            "filename":
                document["filename"],

            "file_type":
                "csv",

            "rows":
                rows

        }


    # XLSX
    if filename.endswith(".xlsx"):

        workbook = openpyxl.load_workbook(
            file_path,
            read_only=True,
            data_only=True
        )

        sheets = {}

        for sheet_name in workbook.sheetnames:

            sheet = workbook[
                sheet_name
            ]

            rows = []

            for row in sheet.iter_rows(
                values_only=True
            ):

                rows.append(
                    list(row)
                )

            sheets[
                sheet_name
            ] = rows

        workbook.close()

        return {

            "document_id":
                document_id,

            "filename":
                document["filename"],

            "file_type":
                "xlsx",

            "sheets":
                sheets

        }


    raise HTTPException(
        status_code=400,
        detail=
        "Preview currently supports CSV and XLSX files"
    )


# =====================================================
# UPDATE DOCUMENT
# =====================================================

class DocumentUpdate(BaseModel):

    filename: str | None = None

    status: str | None = None


@router.put("/{document_id}")
def update_document(
    document_id: str,
    document: DocumentUpdate
):

    update_data = {
        key: value
        for key, value in
        document.model_dump().items()
        if value is not None
    }

    result = db.documents.update_one(
        {
            "id": document_id
        },
        {
            "$set":
                update_data
        }
    )

    if result.matched_count == 0:

        raise HTTPException(
            status_code=404,
            detail="Document not found"
        )

    updated = db.documents.find_one(
        {
            "id": document_id
        },
        {
            "_id": 0
        }
    )

    return {

        "message":
            "Document updated successfully",

        "document":
            updated

    }


# =====================================================
# DELETE DOCUMENT
# =====================================================

@router.delete("/{document_id}")
def delete_document(
    document_id: str
):

    document = db.documents.find_one(
        {
            "id": document_id
        }
    )

    if document is None:

        raise HTTPException(
            status_code=404,
            detail="Document not found"
        )

    file_path = document.get(
        "file_path"
    )

    if file_path:

        file_path = Path(
            file_path
        )

        if file_path.is_file():

            try:
                file_path.unlink()

            except Exception as e:

                print(
                    "Could not delete physical file:",
                    e
                )

    db.documents.delete_one(
        {
            "id": document_id
        }
    )

    return {

        "message":
            "Document deleted successfully"

    }